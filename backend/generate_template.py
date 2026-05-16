"""Generate exam template v4.0.6"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
ws = wb.active
ws.title = '题目模板'

header_font = Font(bold=True, size=11, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
body_align = Alignment(horizontal='left', vertical='top', wrap_text=True)
note_font = Font(size=10, color='333333', bold=True)

ws.merge_cells("A1:M1")
cell = ws['A1']
cell.value = (
    '填写说明：\n'
    + '1. 题型列：使用下拉菜单选择（单选题/多选题/判断题/简答题/综合题）。\n'
    + '2. 题干列：填写题目内容。图片在备注列标注文件名（多张用;隔开），图片于附件上传。\n'
    + '3. 选项列(A-D)：选择项含图片时，在对应选项单元格填写图片文件名（如 chart.png）。\n'
    + '4. 分值列：必填，正整数。\n'
    + '5. 附件列：下拉选择"是"表示该题需要附件，留空则不需要。\n'
    + '6. 参考答案：单选题填大写字母（A），多选题填字母组合（ABD），判断题填"正确/错误"。\n'
    + '7. 备注列：标注题干图片文件名。\n'
    + '8. Stata列：下拉选择[是/否]，默认[是]。\n'
    + '9. AI列：下拉选择[是/否]，默认[是]。\n'
    + '10. 数据列：填写该题需要的数据文件名(.dta/.xlsx)，多个用;隔开。数据文件通过附件管理上传。'
)
cell.font = note_font
cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
ws.row_dimensions[1].height = 200

# New header layout: columns B(序号) removed, I=附件 added, K=数据 added
# A=题型, B=题干, C=选项A, D=选项B, E=选项C, F=选项D, G=参考答案, H=分值, I=附件, J=备注, K=数据, L=Stata, M=AI
headers = ['题型', '题干', '选项A', '选项B', '选项C', '选项D', '参考答案', '分值', '附件', '备注', '数据', 'Stata', 'AI']
for col_idx, h in enumerate(headers, 1):
    c = ws.cell(row=2, column=col_idx, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = header_align

widths = {'A': 14, 'B': 40, 'C': 18, 'D': 18, 'E': 18, 'F': 18,
          'G': 12, 'H': 10, 'I': 8, 'J': 35, 'K': 12, 'L': 8, 'M': 8}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

# Examples: [题型, 题干, A, B, C, D, 参考答案, 分值, 附件, 备注, 数据, Stata, AI]
examples = [
    ('单选题', '以下哪个是Python的特点？', '编译型', '解释型', '汇编', '机器码', 'B', 5, '', '', '', '是', '是'),
    ('多选题', '以下哪些是Web前端技术？', 'HTML', 'CSS', 'Python', 'JavaScript', 'ABD', 5, '', '', '', '是', '是'),
    ('判断题', 'Python是一种面向对象的语言。', '', '', '', '', '正确', 5, '', '', '', '是', '否'),
    ('简答题', '简述Python中的装饰器是什么。', '', '', '', '', '', 10, '', '', '', '是', '是'),
    ('综合题', '阅读以下材料并回答问题...', '', '', '', '', '', 20, '是', 'chart.png', 'data1.dta', '否', '是'),
]
for row_idx, row_data in enumerate(examples, 3):
    for col_idx, value in enumerate(row_data, 1):
        c = ws.cell(row=row_idx, column=col_idx, value=value)
        c.alignment = body_align

# Dropdown for 题型
dv_type = DataValidation(type='list', formula1='"单选题,多选题,判断题,简答题,综合题"',
                          allow_blank=True, showDropDown=False)
dv_type.error = '请从下拉菜单选择题型'
dv_type.errorTitle = '题型选择错误'
dv_type.prompt = '请选择题型'
dv_type.promptTitle = '题型'
ws.add_data_validation(dv_type)
for r in range(3, 105):
    dv_type.add(ws.cell(row=r, column=1))

# Dropdown for 附件
dv_attach = DataValidation(type='list', formula1='"是"',
                            allow_blank=True, showDropDown=False)
dv_attach.error = '请选择"是"或留空'
dv_attach.errorTitle = '附件选择错误'
dv_attach.prompt = '选择"是"或留空'
dv_attach.promptTitle = '附件'
ws.add_data_validation(dv_attach)
for r in range(3, 105):
    dv_attach.add(ws.cell(row=r, column=9))

# Dropdown for Stata
dv_stata = DataValidation(type='list', formula1='"是,否"',
                           allow_blank=True, showDropDown=False)
dv_stata.error = '请选择是/否'
dv_stata.errorTitle = 'Stata选择错误'
ws.add_data_validation(dv_stata)
for r in range(3, 105):
    dv_stata.add(ws.cell(row=r, column=12))

# Dropdown for AI
dv_ai = DataValidation(type='list', formula1='"是,否"',
                        allow_blank=True, showDropDown=False)
dv_ai.error = '请选择是/否'
dv_ai.errorTitle = 'AI选择错误'
ws.add_data_validation(dv_ai)
for r in range(3, 105):
    dv_ai.add(ws.cell(row=r, column=13))

ws.freeze_panes = 'A3'

path = os.path.join(os.path.dirname(__file__), 'exam_template.xlsx')
wb.save(path)
print(f'Template saved: {path}')
